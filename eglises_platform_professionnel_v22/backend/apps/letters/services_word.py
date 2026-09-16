from io import BytesIO
from docx import Document
from django.core.files.base import ContentFile
from django.utils import timezone


def donnees_membre(membre, eglise_destinataire=None):
    return {
        "{{NOM}}": (membre.nom or ""),
        "{{PRENOM}}": (membre.prenom or ""),
        "{{NOM_COMPLET}}": f"{membre.prenom} {membre.nom}".strip(),
        "{{IDENTIFIANT}}": membre.identifiant or "",
        "{{TELEPHONE}}": membre.telephone or "",
        "{{FONCTION}}": membre.get_fonction_eglise_display() if getattr(membre, "fonction_eglise", None) else "Membre",
        "{{EGLISE}}": membre.eglise.nom if getattr(membre, "eglise", None) else "",
        "{{CODE_EGLISE}}": membre.eglise.code if getattr(membre, "eglise", None) else "",
        "{{EGLISE_DESTINATION}}": eglise_destinataire.nom if eglise_destinataire else "",
        "{{CODE_EGLISE_DESTINATION}}": eglise_destinataire.code if eglise_destinataire else "",
        "{{DATE}}": timezone.localdate().strftime("%d/%m/%Y"),
    }


def _remplacer_paragraphe(paragraph, mapping):
    # Conserve la mise en forme du premier run autant que possible.
    texte = paragraph.text
    nouveau = texte
    for ancien, valeur in mapping.items():
        nouveau = nouveau.replace(ancien, valeur)
    if nouveau != texte:
        if paragraph.runs:
            paragraph.runs[0].text = nouveau
            for run in paragraph.runs[1:]:
                run.text = ""
        else:
            paragraph.add_run(nouveau)


def personaliser_word(uploaded_file, membre, eglise_destinataire=None):
    uploaded_file.seek(0)
    doc = Document(BytesIO(uploaded_file.read()))
    mapping = donnees_membre(membre, eglise_destinataire)
    for paragraph in doc.paragraphs:
        _remplacer_paragraphe(paragraph, mapping)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    _remplacer_paragraphe(paragraph, mapping)
    output = BytesIO()
    doc.save(output)
    output.seek(0)
    return ContentFile(output.read(), name=f"recommandation-{membre.identifiant}.docx")
