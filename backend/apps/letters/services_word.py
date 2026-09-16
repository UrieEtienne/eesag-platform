from io import BytesIO
from docx import Document
from django.core.files.base import ContentFile
from django.utils import timezone


def donnees_membre(membre, eglise_destinataire=None):
    origine = getattr(membre, "eglise", None)
    departement = getattr(membre, "departement", None)
    role_eglise = getattr(membre, "role_eglise", None)
    return {
        "{{NOM}}": membre.nom or "",
        "{{PRENOM}}": membre.prenom or "",
        "{{NOM_COMPLET}}": f"{membre.prenom} {membre.nom}".strip(),
        "{{IDENTIFIANT}}": membre.identifiant or "",
        "{{TELEPHONE}}": membre.telephone or "",
        "{{EMAIL}}": membre.email or "",
        "{{SEXE}}": membre.get_sexe_display() if getattr(membre, "sexe", None) else "",
        "{{DATE_NAISSANCE}}": membre.date_naissance.strftime("%d/%m/%Y") if membre.date_naissance else "",
        "{{NATIONALITE}}": membre.get_nationalite_display() if getattr(membre, "nationalite", None) else "",
        "{{FONCTION}}": membre.get_fonction_eglise_display() if getattr(membre, "fonction_eglise", None) else "Membre",
        "{{ROLE_EGLISE}}": role_eglise.nom if role_eglise else "",
        "{{DEPARTEMENT}}": departement.nom if departement else "",
        "{{EGLISE}}": origine.nom if origine else "",
        "{{CODE_EGLISE}}": origine.code if origine else "",
        "{{ADRESSE_EGLISE}}": origine.adresse_precise if origine else "",
        "{{TELEPHONE_EGLISE}}": origine.telephone if origine else "",
        "{{EMAIL_EGLISE}}": origine.email if origine else "",
        "{{REGION_EGLISE}}": getattr(getattr(origine, "region", None), "nom", "") if origine else "",
        "{{PREFECTURE_EGLISE}}": getattr(getattr(origine, "prefecture", None), "nom", "") if origine else "",
        "{{DISTRICT_EGLISE}}": getattr(getattr(origine, "district", None), "nom", "") if origine else "",
        "{{COMMUNE_EGLISE}}": getattr(getattr(origine, "commune", None), "nom", "") if origine and origine.commune_id else "",
        "{{EGLISE_DESTINATION}}": eglise_destinataire.nom if eglise_destinataire else "",
        "{{CODE_EGLISE_DESTINATION}}": eglise_destinataire.code if eglise_destinataire else "",
        "{{REGION_DESTINATION}}": getattr(getattr(eglise_destinataire, "region", None), "nom", "") if eglise_destinataire else "",
        "{{PREFECTURE_DESTINATION}}": getattr(getattr(eglise_destinataire, "prefecture", None), "nom", "") if eglise_destinataire else "",
        "{{DISTRICT_DESTINATION}}": getattr(getattr(eglise_destinataire, "district", None), "nom", "") if eglise_destinataire else "",
        "{{COMMUNE_DESTINATION}}": getattr(getattr(eglise_destinataire, "commune", None), "nom", "") if eglise_destinataire and eglise_destinataire.commune_id else "",
        "{{DATE}}": timezone.localdate().strftime("%d/%m/%Y"),
    }


def _remplacer_paragraphe(paragraph, mapping):
    # Conserve la mise en forme du premier run autant que possible.
    texte = paragraph.text
    nouveau = texte
    for ancien, valeur in mapping.items():
        nouveau = nouveau.replace(ancien, valeur)
    if nouveau != texte:
        if paragraph.runs:
            paragraph.runs[0].text = nouveau
            for run in paragraph.runs[1:]:
                run.text = ""
        else:
            paragraph.add_run(nouveau)


def personaliser_word(uploaded_file, membre, eglise_destinataire=None):
    uploaded_file.seek(0)
    doc = Document(BytesIO(uploaded_file.read()))
    mapping = donnees_membre(membre, eglise_destinataire)
    for paragraph in doc.paragraphs:
        _remplacer_paragraphe(paragraph, mapping)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    _remplacer_paragraphe(paragraph, mapping)
    output = BytesIO()
    doc.save(output)
    output.seek(0)
    return ContentFile(output.read(), name=f"recommandation-{membre.identifiant}.docx")
